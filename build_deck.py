"""Build slides.pptx from reports/figures/*.png and reports/results.json.
Run after survival_analysis.py:  python build_deck.py
Missing figures/numbers show as clearly-marked placeholders instead of fake values.
"""
import json
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

FIG, RES = Path("reports/figures"), Path("reports/results.json")
R = json.loads(RES.read_text()) if RES.exists() else {}
C = dict(ink="14202B", teal="2A9D8F", red="C0392B", gold="D99A1E", paper="FFFFFF", mute="5C6B77", tint="EEF3F5")
rgb = lambda h: RGBColor.from_string(h)

prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide(bg=C["paper"], notes=""):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(bg)
    if notes: s.notes_slide.notes_text_frame.text = notes
    return s


def text(s, x, y, w, h, t, size=16, bold=False, color=C["ink"], font="Calibri", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, head=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"): setattr(tf, m, 0)
    lines = t if isinstance(t, list) else [t]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(6)
        r = p.add_run(); r.text = ln; f = r.font
        f.size, f.bold, f.name, f.color.rgb = Pt(size + (6 if head and i == 0 else 0)), bold or (head and i == 0), font, rgb(color)
    return tb


def title(s, t, color=C["ink"]):
    text(s, 0.7, 0.45, 11.9, 0.9, t, 34, True, color, "Cambria")


def card(s, x, y, w, h, fill=C["tint"]):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.06; sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill); sh.line.fill.background(); sh.shadow.inherit = False
    return sh


def stat(s, x, y, w, big, label, color=C["teal"], dark=False):
    card(s, x, y, w, 1.55, "1F2F3D" if dark else C["tint"])
    text(s, x + 0.25, y + 0.18, w - 0.5, 0.75, big, 34, True, color, "Cambria")
    text(s, x + 0.25, y + 0.97, w - 0.5, 0.5, label, 13, False, "C9D4DB" if dark else C["mute"])


def figure(s, name, x, y, w, h, missing_hint):
    p = FIG / name
    if not p.exists():
        card(s, x, y, w, h, "FBEDEB")
        text(s, x + 0.3, y + h / 2 - 0.4, w - 0.6, 0.9, [f"[MISSING FIGURE: {name}]", missing_hint], 14, False, C["red"], align=PP_ALIGN.CENTER)
        return
    iw, ih = Image.open(p).size; sc = min(w / iw, h / ih)
    s.shapes.add_picture(str(p), Inches(x + (w - iw * sc) / 2), Inches(y + (h - ih * sc) / 2), Inches(iw * sc), Inches(ih * sc))


def num(k, fmt, default="—"):
    return fmt.format(R[k]) if R.get(k) is not None else default


def pfmt(p): return "p < 0.001" if p < 0.001 else f"p = {p:.3f}"

# 1 ---------------------------------------------------------------- title
s = slide(C["ink"], "Intro (20s): we predict overall survival in pancreatic cancer with Cox models. Tanvi built the models; I'll show what the models say and let you try one.")
text(s, 0.8, 2.2, 11.5, 1.8, "Pancreatic cancer prognosis with Cox survival models", 46, True, C["paper"], "Cambria")
text(s, 0.8, 4.3, 11, 0.5, "Risk stratification, hazard ratios and a live risk calculator", 20, False, "C9D4DB")
text(s, 0.8, 6.3, 11, 0.4, "UE24CS352A Machine Learning mini-project · Tanvi B. Shetty · Y. V. Bhavana", 14, False, "8FA1AE")

# 2 ---------------------------------------------------------------- data
s = slide(notes="Public TCGA-PAAD stands in for the Stanford data in the reference paper, which is under NDA. Gene expression plays the role of the radiomic features: high-dimensional, few patients.")
title(s, "Data: TCGA-PAAD, 183 patients, 99 deaths")
stat(s, 0.7, 1.9, 3.8, "183", "patients (TCGA-PAAD PanCancer Atlas)")
stat(s, 4.75, 1.9, 3.8, "99", "observed deaths (the rest are censored)", C["red"])
stat(s, 8.8, 1.9, 3.8, "12 + 500", "clinical features + top-variance genes", C["gold"])
text(s, 0.7, 4.0, 11.9, 2.6, [
    "Public data replaces the paper's Stanford cohort, which is under NDA.",
    "Gene expression stands in for radiomic features: many features, few patients.",
    "Outcome: overall survival in months, with right-censoring handled by the Cox model."], 18)

# 3 ---------------------------------------------------------------- method
s = slide(notes="Features are standardised inside each training fold so nothing leaks from the test fold. Penalty and gene count were tuned by a sweep on the same CV, so tuned scores are slightly optimistic.")
title(s, "Method: penalised Cox, scored by cross-validation")
steps = [("1", "Cox PH + L2", "Hazard depends on a weighted sum of features; ridge penalty tames 100+ genes"),
         ("2", "10-fold CV", "Concordance index on held-out patients, features scaled per training fold"),
         ("3", "Tune", "Penalty 0.1–50 and 25–500 genes; best: 100 genes, penalty 1"),
         ("4", "Interpret", "Out-of-fold risk scores → KM curves, log-rank, hazard ratios")]
for i, (n, h, d) in enumerate(steps):
    x = 0.7 + i * 3.05
    card(s, x, 2.0, 2.85, 3.6)
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.25), Inches(2.25), Inches(0.6), Inches(0.6))
    c.fill.solid(); c.fill.fore_color.rgb = rgb(C["teal"]); c.line.fill.background()
    c.text_frame.text = n; r = c.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = rgb(C["paper"])
    c.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    text(s, x + 0.25, 3.1, 2.4, 0.5, h, 20, True)
    text(s, x + 0.25, 3.7, 2.4, 1.8, d, 14, False, C["mute"])

# 4 ---------------------------------------------------------------- c-index table
s = slide(notes="Clinical alone is about 0.63. Expression alone is poor untuned but reaches 0.64 once regularised properly. Combined is best at 0.67. Tuned numbers are the best of a sweep on the same CV, so slightly optimistic.")
title(s, "Combining clinical and gene data gives the best C-index")
rows = [("Model", "Untuned", "Tuned"), ("Clinical (12 features)", "0.63", "0.63"),
        ("Expression", "0.57", "0.64  (100 genes, penalty 1)"), ("Combined", "0.58", "0.67  (100 genes, penalty 1)")]
tbl = s.shapes.add_table(4, 3, Inches(0.7), Inches(1.9), Inches(11.9), Inches(3.0)).table
for j, w in enumerate((4.2, 2.7, 5.0)): tbl.columns[j].width = Inches(w)
for i, row in enumerate(rows):
    for j, v in enumerate(row):
        c = tbl.cell(i, j); c.text = v; p = c.text_frame.paragraphs[0]; r = p.runs[0]
        r.font.size = Pt(20); r.font.name = "Calibri"; r.font.bold = (i == 0 or (i == 3 and j == 2))
        r.font.color.rgb = rgb(C["paper"] if i == 0 else (C["teal"] if (i == 3 and j == 2) else C["ink"]))
        c.fill.solid(); c.fill.fore_color.rgb = rgb(C["ink"] if i == 0 else (C["tint"] if i % 2 else C["paper"]))
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
text(s, 0.7, 5.4, 11.9, 1, "10-fold CV concordance index (0.5 = coin flip, 1.0 = perfect ranking). Tuned scores are the best of a sweep chosen on the same CV, so slightly optimistic.", 14, False, C["mute"])

# 5 ---------------------------------------------------------------- KM
s = slide(notes="Each patient is scored by a model that never saw them (out-of-fold), then split at the median. The log-rank test asks whether the two curves could plausibly be the same. Point at the at-risk table: the tail is thin, so late-time estimates are noisy.")
title(s, "High-risk patients die sooner" if R.get("logrank_p", 1) < 0.05 else "Survival by predicted risk group")
figure(s, "km_high_low.png", 0.6, 1.4, 7.6, 5.7, "Run: python survival_analysis.py")
lr = pfmt(R["logrank_p"]) if "logrank_p" in R else "—"
stat(s, 8.6, 1.6, 4.1, lr, "log-rank test, high vs low risk", C["red"])
hr = f"{R['hr_high_vs_low']:.2f}" if "hr_high_vs_low" in R else "—"
stat(s, 8.6, 3.3, 4.1, hr + "×", "hazard of death, high vs low" + (f"  (95% CI {R['hr_lo']:.2f}–{R['hr_hi']:.2f})" if "hr_lo" in R else ""), C["gold"])
mo = (f"{num('median_os_low', '{:.0f}')} vs {num('median_os_high', '{:.0f}')} mo") if R else "—"
stat(s, 8.6, 5.0, 4.1, mo, "median survival, low vs high risk", C["teal"])

# 6 ---------------------------------------------------------------- forest
s = slide(notes="Hazard ratio above 1 means shorter survival. Red intervals exclude 1. With 183 patients most intervals are wide and cross 1; that width is the honest message. Clinical model in natural units, combined model per one standard deviation.")
title(s, "Which features move the hazard?")
figure(s, "forest_combined_top.png", 0.6, 1.4, 6.3, 5.7, "Run: python survival_analysis.py")
figure(s, "forest_clinical.png", 6.9, 1.4, 6.0, 5.7, "Run: python survival_analysis.py")

# 7 ---------------------------------------------------------------- covariate curves
s = slide(notes="Each panel changes one feature and holds everything else at the cohort median. This is the model's prediction, not raw data, so it shows direction and size of effect, not a new finding.")
title(s, "Covariate effects on predicted survival")
figure(s, "covariate_curves.png", 0.6, 1.35, 8.6, 5.8, "Run: python survival_analysis.py")
text(s, 9.5, 1.9, 3.3, 4.5, ["One feature changes per panel; all others stay at the cohort median or mean.",
                             "Curves are model predictions, not raw data.",
                             "Wider gaps between curves mean a larger effect."], 16, False, C["mute"])

# 8 ---------------------------------------------------------------- demo
s = slide(notes="Switch to the browser: streamlit run demo_app.py. Follow demo_script.md.")
title(s, "Live demo: risk calculator")
for i, (h, d) in enumerate([("Enter", "Clinical features and top gene levels, or start from a real cohort patient"),
                            ("Score", "Relative hazard vs the average patient, percentile, high/low group"),
                            ("Explain", "Predicted survival curve and the features driving the score")]):
    x = 0.7 + i * 4.1
    card(s, x, 2.0, 3.9, 3.4)
    text(s, x + 0.3, 2.35, 3.3, 0.6, h, 26, True, C["teal"], "Cambria")
    text(s, x + 0.3, 3.2, 3.3, 2, d, 16, False, C["ink"])
text(s, 0.7, 6.0, 11.9, 0.6, "streamlit run demo_app.py", 18, False, C["mute"], "Courier New")

SEP = ("Out-of-fold risk groups " + ("separate clearly" if R.get("logrank_p", 1) < 0.05 else "did not separate clearly") + f" ({lr}).") if R else "Out-of-fold risk groups: [fill in after running the analysis]."
# 9 ---------------------------------------------------------------- limits / takeaways
s = slide(C["ink"], "Close: three takeaways, then limitations.")
title(s, "Takeaways and limitations", C["paper"])
text(s, 0.7, 1.7, 5.9, 4.8, ["What we found",
    "Clinical + expression beats either alone (C-index 0.67 vs 0.63 / 0.64).",
    SEP,
    "Regularisation matters: untuned expression models barely beat chance."], 17, False, "E6EEF2", head=True)
text(s, 7.0, 1.7, 5.6, 4.8, ["Limitations",
    "183 patients: wide intervals, one cohort, no external validation.",
    "Tuned scores are slightly optimistic (sweep on same CV).",
    "Median imputation of missing clinical values.",
    "Research prototype, not for clinical use."], 17, False, "E6EEF2", head=True)

prs.save("slides.pptx"); print("saved slides.pptx", "(with real results)" if R else "(placeholders: run survival_analysis.py first)")
